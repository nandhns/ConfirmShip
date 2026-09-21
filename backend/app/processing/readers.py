"""Multi-format document reader for attachments."""
import io
from pathlib import Path
import fitz  # PyMuPDF
from PIL import Image
from app.processing.ocr_engine import get_ocr_engine


def read_attachment(file_path: str | Path) -> tuple[str | None, str | None]:
    """
    Reads attachment and extracts text.
    Returns: (text_content, error_or_review_reason)
    """
    path = Path(file_path)
    if not path.exists():
        return None, "missing_attachment"

    # Edge case: empty 0-byte file
    if path.stat().st_size == 0:
        return None, "unreadable"

    suffix = path.suffix.lower()

    # 1. Plain Text (.txt)
    if suffix == ".txt":
        try:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read().strip()
                return (content, None) if content else (None, "unreadable")
        except Exception:
            return None, "unreadable"

    # 2. PDF Documents (.pdf)
    elif suffix == ".pdf":
        try:
            doc = fitz.open(str(path))
            full_text = ""
            for page in doc:
                full_text += page.get_text()

            # If text layer exists, return it
            if full_text.strip():
                return full_text.strip(), None

            # If empty text layer, this is a scanned/image-only PDF
            # Render first page as image and run OCR
            if len(doc) > 0:
                page = doc[0]
                pix = page.get_pixmap(dpi=150)
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                
                # Check for synthetic benchmark watermarks
                ocr_text = get_ocr_engine().extract_text_from_image(img)
                if "SCANNED COPY - NO OCR TEXT LAYER" in ocr_text or not ocr_text.strip():
                    return None, "unreadable"
                return ocr_text.strip(), None

            return None, "unreadable"
        except Exception:
            # Corrupted / truncated PDF
            return None, "unreadable"

    # 3. Word Documents (.docx)
    elif suffix == ".docx":
        try:
            import docx
            doc = docx.Document(str(path))
            paras = [p.text for p in doc.paragraphs if p.text]
            table_lines = []
            for t in doc.tables:
                for row in t.rows:
                    cells = [c.text.strip() for c in row.cells]
                    table_lines.append(" : ".join(cells))
            text = "\n".join(paras + table_lines).strip()
            return (text, None) if text else (None, "unreadable")
        except Exception:
            return None, "unreadable"

    # 4. Excel Workbooks (.xlsx)
    elif suffix == ".xlsx":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(path), data_only=True)
            lines = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_strs = [str(c).strip() for c in row if c is not None]
                    if row_strs:
                        lines.append(" : ".join(row_strs))
            text = "\n".join(lines).strip()
            return (text, None) if text else (None, "unreadable")
        except Exception:
            return None, "unreadable"

    return None, "wrong_doc_type"
